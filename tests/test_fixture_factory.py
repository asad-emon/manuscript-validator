"""Task 3 exit criteria: the fixture factory builds a structurally-compliant
manuscript, and every registered `violating(rule_id)` mutation changes exactly
the snapshot field(s) declared for that rule and nothing else.

No dependency on `manuscript_validator.parser` -- it does not exist yet
(Task 4). Structure is checked directly against python-docx/lxml, which is
what makes this test independent of the (unbuilt) AST.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from fixtures.factory import (
    MUTATION_TARGETS,
    RULE_IDS,
    _before,
    build_compliant,
    diff,
    snapshot,
    violating,
)


def _roundtrip(doc: Document) -> Document:
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Document(buf)


def test_compliant_covers_all_required_structure() -> None:
    doc, handles = build_compliant()

    heading_texts = [p.text for p in doc.paragraphs if p.runs and p.runs[0].font.bold]
    assert any(text.startswith("INTRODUCTION") for text in heading_texts)

    assert len(doc.tables) >= 2
    assert len(doc.inline_shapes) >= 2  # two figures inserted via Run.add_picture
    assert len(handles.reference_runs) >= 8

    # Superscript affiliation numeral and superscript bracket citation.
    assert handles.author_super_run.font.superscript is True
    assert handles.citation_super_run.font.superscript is True
    assert handles.citation_super_run.text == "[1]"

    # Roman table numerals, Arabic figure numerals.
    assert "Table I." in handles.table1.caption_run.text
    assert "Table II." in handles.table2.caption_run.text
    assert "Figure 1." in handles.figure1.caption_run.text
    assert "Figure 2." in handles.figure2.caption_run.text

    # Caption position: table caption above the table, figure caption below the image.
    assert _before(*handles.order_pairs["table1_caption_vs_table"]) is True
    assert _before(*handles.order_pairs["figure1_caption_vs_image"]) is False


def test_compliant_round_trips_without_corruption() -> None:
    doc, _ = build_compliant()
    reopened = _roundtrip(doc)
    assert len(reopened.tables) == 2
    assert len(reopened.paragraphs) > 0


def test_rule_inventory_has_no_duplicates_and_matches_targets() -> None:
    assert len(RULE_IDS) == len(set(RULE_IDS))
    assert set(RULE_IDS) == set(MUTATION_TARGETS)
    assert len(RULE_IDS) == 40


@pytest.mark.parametrize("rule_id", RULE_IDS)
def test_violating_changes_exactly_the_declared_fields(rule_id: str) -> None:
    _, baseline_handles = build_compliant()
    baseline = snapshot(baseline_handles)

    doc, violating_handles = violating(rule_id)
    changed = diff(baseline, snapshot(violating_handles))

    assert changed == MUTATION_TARGETS[rule_id]
    # And the mutated document is still well-formed enough to reopen.
    _roundtrip(doc)
