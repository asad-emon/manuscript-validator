"""Task 14 exit criteria (report/override-panel half): `ReportView` renders
every violation, and `SectionOverridePanel` only reports the paragraphs a
human actually changed -- not every row's current (possibly unchanged) value.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from PySide6.QtWidgets import QComboBox

from fixtures.factory import build_compliant, violating
from manuscript_validator.models.enums import Section
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.segmenter import segment
from manuscript_validator.ui.report_view import ReportView, SectionOverridePanel

pytestmark = pytest.mark.gui

RULESET = load_ruleset("journal_v1")


def _ast_for(doc: Document):
    buf = BytesIO()
    doc.save(buf)
    source_bytes = buf.getvalue()
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segment(ast, RULESET)
    return ast


def test_report_view_renders_one_row_per_violation(qtbot) -> None:
    doc, _handles = violating("title-size")
    ast = _ast_for(doc)
    segmentation = segment(ast, RULESET)
    violations = rules_engine.validate(ast, RULESET, segmentation)

    from manuscript_validator.report import build_report

    report = build_report(ast.document_id, RULESET.ruleset_version, violations, segmentation)

    view = ReportView()
    qtbot.addWidget(view)
    view.set_report(report)

    assert view.rowCount() == len(report.violations)
    rule_ids = {view.item(row, 0).text() for row in range(view.rowCount())}
    assert "title-size" in rule_ids


def test_override_panel_lists_every_paragraph_with_its_detected_section(qtbot) -> None:
    doc, _handles = build_compliant()
    ast = _ast_for(doc)

    panel = SectionOverridePanel()
    qtbot.addWidget(panel)
    panel.set_ast(ast)

    assert panel._table.rowCount() == len(ast.paragraphs)
    title_row_combo = panel._table.cellWidget(0, 2)
    assert isinstance(title_row_combo, QComboBox)
    assert title_row_combo.currentData() == "title"


def test_override_panel_only_emits_changed_rows(qtbot) -> None:
    doc, _handles = build_compliant()
    ast = _ast_for(doc)

    panel = SectionOverridePanel()
    qtbot.addWidget(panel)
    panel.set_ast(ast)

    # Row 0 is the title paragraph; change it to "author" without touching
    # any other row.
    combo = panel._table.cellWidget(0, 2)
    index = combo.findData(Section.AUTHOR.value)
    combo.setCurrentIndex(index)

    with qtbot.waitSignal(panel.overrides_applied, timeout=1000) as blocker:
        panel._on_apply()

    overrides = blocker.args[0]
    assert overrides == {ast.paragraphs[0].id: "author"}


def test_override_panel_emits_nothing_when_no_row_changed(qtbot) -> None:
    doc, _handles = build_compliant()
    ast = _ast_for(doc)

    panel = SectionOverridePanel()
    qtbot.addWidget(panel)
    panel.set_ast(ast)

    with qtbot.waitSignal(panel.overrides_applied, timeout=1000) as blocker:
        panel._on_apply()

    assert blocker.args[0] == {}
