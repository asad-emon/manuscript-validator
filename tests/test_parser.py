"""Task 4 exit criteria: effective formatting resolution across all four
inheritance sources (including theme), canonical text extraction, id
stability across independently-parsed clones, and a full AST build over
Task 3's compliant and violating fixtures.
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from itertools import pairwise

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from lxml import etree

from fixtures.factory import RULE_IDS, build_compliant, violating
from manuscript_validator.models.enums import CaptionPosition, FormattingSource, NumberingStyle
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.parser.binding import build_element_index
from manuscript_validator.parser.captions import (
    classify_numbering,
    find_citing_paragraph_ids,
    find_figure_caption,
    find_table_caption,
)
from manuscript_validator.parser.effective import (
    _MAX_STYLE_DEPTH,
    EffectiveFormattingResolver,
    _style_chain,
    is_uppercase_literal,
    paragraph_text,
    run_text,
)
from manuscript_validator.parser.ids import ParagraphItem, walk_document
from manuscript_validator.parser.shapes import find_figure_paragraphs


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# paragraph_text / run_text / is_uppercase_literal
# --------------------------------------------------------------------------


def test_paragraph_text_normalizes_tab_break_and_nbsp() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("A")
    tab_run = p.add_run()
    tab_run.add_tab()
    p.add_run("B C")
    assert paragraph_text(p._p) == "A\tB C"


def test_paragraph_text_includes_hyperlink_and_ins_but_not_rprchange() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p_el = p._p
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    hyperlink = etree.SubElement(p_el, f"{{{ns}}}hyperlink")
    run = etree.SubElement(hyperlink, f"{{{ns}}}r")
    text_el = etree.SubElement(run, f"{{{ns}}}t")
    text_el.text = "linked"

    ins = etree.SubElement(p_el, f"{{{ns}}}ins")
    ins_run = etree.SubElement(ins, f"{{{ns}}}r")
    ins_text = etree.SubElement(ins_run, f"{{{ns}}}t")
    ins_text.text = "inserted"

    # A run nested inside another run's rPrChange snapshot must be excluded.
    plain_run = etree.SubElement(p_el, f"{{{ns}}}r")
    rpr = etree.SubElement(plain_run, f"{{{ns}}}rPr")
    rpr_change = etree.SubElement(rpr, f"{{{ns}}}rPrChange")
    old_rpr = etree.SubElement(rpr_change, f"{{{ns}}}rPr")
    ghost_run = etree.SubElement(old_rpr, f"{{{ns}}}r")
    ghost_text = etree.SubElement(ghost_run, f"{{{ns}}}t")
    ghost_text.text = "GHOST"
    visible_text = etree.SubElement(plain_run, f"{{{ns}}}t")
    visible_text.text = "visible"

    text = paragraph_text(p_el)
    assert text == "linkedinsertedvisible"
    assert "GHOST" not in text


def test_is_uppercase_literal() -> None:
    assert is_uppercase_literal("INTRODUCTION") is True
    assert is_uppercase_literal("Introduction") is False
    assert is_uppercase_literal("123") is False  # no letters at all


def test_run_text_matches_docx_run_text() -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("A B")
    assert run_text(run._r) == "A B"


# --------------------------------------------------------------------------
# EffectiveFormattingResolver: all four inheritance sources
# --------------------------------------------------------------------------


def test_resolves_from_run_level() -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("x")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    resolver = EffectiveFormattingResolver(doc)
    name = resolver.resolve_font_name(run, p)
    size = resolver.resolve_font_size(run, p)
    bold = resolver.resolve_bold(run, p)
    assert (name.value, name.source) == ("Times New Roman", FormattingSource.RUN)
    assert (size.value, size.source) == (14.0, FormattingSource.RUN)
    assert (bold.value, bold.source) == (True, FormattingSource.RUN)


def test_resolves_from_character_style() -> None:
    doc = Document()
    char_style = doc.styles.add_style("Emph", WD_STYLE_TYPE.CHARACTER)
    char_style.font.size = Pt(16)
    char_style.font.bold = True
    p = doc.add_paragraph()
    run = p.add_run("x")
    run.style = char_style
    resolver = EffectiveFormattingResolver(doc)
    size = resolver.resolve_font_size(run, p)
    bold = resolver.resolve_bold(run, p)
    assert (size.value, size.source) == (16.0, FormattingSource.CHARACTER_STYLE)
    assert (bold.value, bold.source) == (True, FormattingSource.CHARACTER_STYLE)


def test_resolves_from_paragraph_style_basedon_chain() -> None:
    doc = Document()
    base = doc.styles.add_style("BaseHeading", WD_STYLE_TYPE.PARAGRAPH)
    base.font.size = Pt(20)
    derived = doc.styles.add_style("DerivedHeading", WD_STYLE_TYPE.PARAGRAPH)
    derived.element.get_or_add_basedOn().val = base.style_id
    p = doc.add_paragraph()
    p.style = derived
    run = p.add_run("x")  # no direct formatting at all
    resolver = EffectiveFormattingResolver(doc)
    size = resolver.resolve_font_size(run, p)
    assert (size.value, size.source) == (20.0, FormattingSource.PARAGRAPH_STYLE)


def test_resolves_from_doc_defaults_and_theme() -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("x")  # nothing set anywhere
    resolver = EffectiveFormattingResolver(doc)
    size = resolver.resolve_font_size(run, p)
    assert size.source == FormattingSource.DOC_DEFAULT
    assert size.value is not None

    # Font name has no docDefaults literal `w:ascii` in the stock template --
    # only a theme reference -- so it must resolve via THEME, matching the
    # theme part's own minor-font typeface (spec-alignment risk 1).
    name = resolver.resolve_font_name(run, p)
    assert name.source == FormattingSource.THEME
    theme_root = etree.fromstring(
        next(
            rel.target_part.blob
            for rel in doc.part.rels.values()
            if rel.reltype.endswith("/theme")
        )
    )
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    minor_latin = theme_root.find(f".//{{{a_ns}}}fontScheme/{{{a_ns}}}minorFont/{{{a_ns}}}latin")
    assert name.value == minor_latin.get("typeface")


def test_bold_tri_state_explicit_false_vs_absent() -> None:
    doc = Document()
    p = doc.add_paragraph()
    explicit_false = p.add_run("a")
    explicit_false.font.bold = False
    absent = p.add_run("b")
    resolver = EffectiveFormattingResolver(doc)
    explicit_resolved = resolver.resolve_bold(explicit_false, p)
    assert (explicit_resolved.value, explicit_resolved.source) == (False, FormattingSource.RUN)
    # Absent everywhere in the chain resolves to False (Word's implicit
    # default), not None -- bold/italic/superscript are never ambiguous.
    resolved = resolver.resolve_bold(absent, p)
    assert resolved.value is False
    assert resolved.source == FormattingSource.UNRESOLVED


def test_style_chain_is_cycle_guarded_and_depth_capped() -> None:
    doc = Document()
    a = doc.styles.add_style("A", WD_STYLE_TYPE.PARAGRAPH)
    b = doc.styles.add_style("B", WD_STYLE_TYPE.PARAGRAPH)
    a.element.get_or_add_basedOn().val = b.style_id
    b.element.get_or_add_basedOn().val = a.style_id  # cycle
    assert len(list(_style_chain(a.element))) <= _MAX_STYLE_DEPTH

    styles = [doc.styles.add_style(f"Chain{i}", WD_STYLE_TYPE.PARAGRAPH) for i in range(30)]
    for prev, nxt in pairwise(styles):
        nxt.element.get_or_add_basedOn().val = prev.style_id
    chain = list(_style_chain(styles[-1].element))
    assert len(chain) == _MAX_STYLE_DEPTH


# --------------------------------------------------------------------------
# ids.walk_document / iter_indexed_runs
# --------------------------------------------------------------------------


def test_walk_document_preserves_paragraph_table_interleaving() -> None:
    doc = Document()
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("after")
    kinds = [
        "paragraph" if isinstance(item, ParagraphItem) else "table" for item in walk_document(doc)
    ]
    assert kinds == ["paragraph", "table", "paragraph", "paragraph"]  # table cell + "after"


def test_walk_document_dedupes_merged_cells() -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "merged"
    items = [item for item in walk_document(doc) if isinstance(item, ParagraphItem)]
    assert len(items) == 3  # one merged cell + two ordinary cells in row 2


# --------------------------------------------------------------------------
# shapes.find_figure_paragraphs
# --------------------------------------------------------------------------


def test_find_figure_paragraphs_dedupes_multiple_drawings_per_paragraph() -> None:
    doc, _handles = build_compliant()
    paragraphs = find_figure_paragraphs(doc)
    assert len(paragraphs) == 2


# --------------------------------------------------------------------------
# captions
# --------------------------------------------------------------------------


def test_classify_numbering() -> None:
    assert classify_numbering("I") == NumberingStyle.ROMAN
    assert classify_numbering("IV") == NumberingStyle.ROMAN
    assert classify_numbering("1") == NumberingStyle.ARABIC
    assert classify_numbering("?") == NumberingStyle.OTHER


def test_caption_pattern_does_not_match_narrative_prose() -> None:
    """A narrative sentence that happens to start with "Table I"/"Figure 1"
    must not be mistaken for the caption -- only "Table I." (period/colon
    right after the number) is a caption. Regression test for the bug this
    exact ambiguity caused during development."""
    _doc, handles = build_compliant()
    match = find_table_caption(handles.table1.table_element)
    assert match is not None
    assert match.text.startswith("Table I.")
    assert "summarizes" not in match.text

    fig_match = find_figure_caption(handles.figure1.image_element)
    assert fig_match is not None
    assert fig_match.text.startswith("Figure 1.")
    assert "illustrates" not in fig_match.text


def test_find_citing_paragraph_ids() -> None:
    _doc, handles = build_compliant()
    match = find_table_caption(handles.table1.table_element)
    assert match is not None
    candidates = [("p1", "Table I summarizes things."), ("p2", "unrelated text")]
    assert find_citing_paragraph_ids(match, candidates) == ["p1"]


# --------------------------------------------------------------------------
# binding.ElementIndex
# --------------------------------------------------------------------------


def test_element_index_ids_stable_across_independent_clones() -> None:
    doc, _ = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    clone_a = Document(BytesIO(source_bytes))
    clone_b = Document(BytesIO(source_bytes))
    index_a = build_element_index(clone_a)
    index_b = build_element_index(clone_b)
    assert set(index_a.paragraphs) == set(index_b.paragraphs)
    assert set(index_a.runs) == set(index_b.runs)
    assert set(index_a.tables) == set(index_b.tables)
    assert index_a.paragraph("p0001").xpath(".//w:t")[0].text == (
        index_b.paragraph("p0001").xpath(".//w:t")[0].text
    )


def test_element_index_survives_reparenting() -> None:
    """Ids resolve to *elements*, not positions: moving a paragraph elsewhere
    in the tree (as a caption-position fix does) must not invalidate its id."""
    doc, _ = build_compliant()
    index = build_element_index(doc)
    caption_id = "p0012"  # table1's caption paragraph, built directly above table1
    caption_el = index.paragraph(caption_id)
    table_el = index.tables["t01"]
    table_el.addnext(caption_el)  # move the caption after the table
    assert index.paragraph(caption_id) is caption_el
    assert list(table_el.getparent()).index(caption_el) > list(
        table_el.getparent()
    ).index(table_el)


def test_element_index_raises_document_error_for_unknown_id() -> None:
    from manuscript_validator.errors import DocumentError

    doc, _ = build_compliant()
    index = build_element_index(doc)
    with pytest.raises(DocumentError):
        index.paragraph("p9999")


# --------------------------------------------------------------------------
# ast_builder.build_ast: the compliant fixture and all 40 violating fixtures
# --------------------------------------------------------------------------


def test_build_ast_populates_every_field_on_the_compliant_fixture() -> None:
    doc, _handles = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)

    assert ast.document_id == hashlib.sha256(source_bytes).hexdigest()
    assert len(ast.paragraphs) > 0
    assert len(ast.tables) == 2
    assert len(ast.figures) == 2

    title = ast.paragraphs[0]
    assert title.runs[0].font_size_pt == 14.0
    assert title.runs[0].bold is True
    assert title.runs[0].font_name == "Times New Roman"

    citation_paragraph = next(p for p in ast.paragraphs if "established evidence" in p.text)
    citation_run = next(r for r in citation_paragraph.runs if r.text == "[1]")
    assert citation_run.superscript is True

    table1 = ast.tables[0]
    assert table1.caption_position == CaptionPosition.ABOVE
    assert table1.numbering_style == NumberingStyle.ROMAN
    assert table1.caption_number == "I"
    assert table1.cell_font_size_pt == 8.0
    assert len(table1.cell_paragraph_ids) == 4
    assert len(table1.referenced_in_paragraph_ids) == 1

    figure1 = ast.figures[0]
    assert figure1.caption_position == CaptionPosition.BELOW
    assert figure1.numbering_style == NumberingStyle.ARABIC
    assert figure1.caption_number == "1"
    assert figure1.paragraph_id is not None
    assert len(figure1.referenced_in_paragraph_ids) == 1


def test_build_ast_never_mutates_the_source_bytes() -> None:
    """FR-9, asserted rather than assumed: parsing is read-only.

    Compares the in-memory `document.xml` tree before/after, not two
    `.save()` outputs -- the zip container embeds a save-time timestamp, so
    even an untouched document re-serializes to different bytes on a second
    call, which would make that comparison a false positive for mutation.
    """
    doc, _ = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    document = Document(BytesIO(source_bytes))
    xml_before = etree.tostring(document.element)
    build_ast(document, source_bytes)
    assert etree.tostring(document.element) == xml_before


def test_parse_parse_id_stability_across_clones() -> None:
    doc, _ = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    ast_a = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    ast_b = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    assert [p.id for p in ast_a.paragraphs] == [p.id for p in ast_b.paragraphs]
    assert [t.id for t in ast_a.tables] == [t.id for t in ast_b.tables]
    assert [f.id for f in ast_a.figures] == [f.id for f in ast_b.figures]


@pytest.mark.parametrize("rule_id", RULE_IDS)
def test_build_ast_handles_every_violating_fixture(rule_id: str) -> None:
    doc, _ = violating(rule_id)
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    assert len(ast.paragraphs) > 0
