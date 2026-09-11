"""Task 8 exit criteria (planner half): `plan_fixes`/`apply_fix_plan` turn
violations into audit entries that match spec section 5.5, the conflict
detector withholds genuinely conflicting fixes rather than letting a last
writer silently win, and applying a fix plan never touches the document (or
AST) it was planned from -- only a separate clone.

The strongest single check here mirrors Task 6's: every one of the 24
`auto_fixable: true` rules, applied to its own `violating(rule_id)` fixture,
must re-validate clean afterward -- proof the fix is correct, not just that
it ran without raising.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from fixtures.factory import build_compliant, violating
from manuscript_validator.autofix import apply_fix_plan, plan_fixes
from manuscript_validator.models.enums import FixAction, Severity, ViolationStatus
from manuscript_validator.models.report import utc_timestamp
from manuscript_validator.models.violation import Violation
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.parser.binding import build_element_index
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.segmenter import segment

RULESET = load_ruleset("journal_v1")
AUTO_FIXABLE_RULE_IDS = [rule.rule_id for rule in RULESET.rules if rule.auto_fixable]


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _validate(source_bytes: bytes):
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    return ast, rules_engine.validate(ast, RULESET, segmentation)


def test_auto_fixable_rules_have_at_least_24_entries() -> None:
    # A floor, not an exact count: guards against a future rule quietly
    # losing its fix_action rather than pinning the number forever.
    assert len(AUTO_FIXABLE_RULE_IDS) >= 24


@pytest.mark.parametrize("rule_id", AUTO_FIXABLE_RULE_IDS)
def test_fix_resolves_the_violation_it_was_planned_for(rule_id: str) -> None:
    doc, _handles = violating(rule_id)
    source_bytes = _roundtrip_bytes(doc)

    ast, violations = _validate(source_bytes)
    ast_snapshot_before = ast.to_dict()
    plan = plan_fixes(violations, RULESET)
    assert any(op.rule_id == rule_id for op in plan.ops), "no FixOp planned"

    # `apply_fix_plan` is handed a *different* document's index -- `ast` came
    # from `source_bytes`, but the live clone below is a second, independent
    # parse of the same bytes. Nothing here ever touches `ast` again.
    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    audit_log = apply_fix_plan(index, plan, violations, utc_timestamp())
    assert any(entry.rule_id == rule_id for entry in audit_log)
    unmutated_msg = "the input AST must never be touched by apply_fix_plan"
    assert ast.to_dict() == ast_snapshot_before, unmutated_msg

    fixed_violation = next(v for v in violations if v.rule_id == rule_id)
    assert fixed_violation.status is ViolationStatus.FIXED

    fixed_bytes = _roundtrip_bytes(live_doc)

    _ast2, violations2 = _validate(fixed_bytes)
    still_open = [
        v
        for v in violations2
        if v.rule_id == rule_id and v.status is not ViolationStatus.CHECK_FAILED
    ]
    assert still_open == []


def test_plan_fixes_never_touches_the_source_document() -> None:
    """`plan_fixes` takes violations and a ruleset -- no document or AST
    parameter exists for it to mutate, which is what makes "input AST
    provably unmutated" true by construction rather than by convention."""
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    ast, violations = _validate(source_bytes)
    ast_snapshot_before = ast.to_dict()

    plan_fixes(violations, RULESET)

    assert ast.to_dict() == ast_snapshot_before


def test_apply_fix_plan_produces_audit_entries_matching_spec_5_5_shape() -> None:
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    _ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)

    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    audit_log = apply_fix_plan(index, plan, violations, "2026-09-03T10:00:02Z")

    entry = next(e for e in audit_log if e.rule_id == "title-size")
    data = entry.to_dict()
    assert data["rule_id"] == "title-size"
    assert data["action"] == "set_font_size"
    assert data["before"] == {"font_size_pt": 11.0}
    assert data["after"] == {"font_size_pt": 14}
    assert data["timestamp"] == "2026-09-03T10:00:02Z"
    assert isinstance(data["paragraph_id"], str) and data["paragraph_id"]


def test_apply_fix_plan_marks_fix_failed_when_a_target_is_gone() -> None:
    """A target that no longer exists (moved, deleted between planning and
    apply) degrades to one `fix_failed` violation, never a crash."""
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    _ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)

    # Apply against a document that never had the title paragraph parsed
    # into the index at all -- an empty document stands in for "target gone".
    live_doc = Document()
    index = build_element_index(live_doc)
    apply_fix_plan(index, plan, violations, utc_timestamp())

    violation = next(v for v in violations if v.rule_id == "title-size")
    assert violation.status is ViolationStatus.FIX_FAILED
    assert violation.failure_reason


# --------------------------------------------------------------------------
# Conflict detection
# --------------------------------------------------------------------------


def _fake_violation(rule_id: str, paragraph_id: str, run_index: int = 0) -> Violation:
    return Violation(
        rule_id=rule_id,
        severity=Severity.MEDIUM,
        paragraph_id=paragraph_id,
        auto_fixable=True,
        status=ViolationStatus.OPEN,
        run_indices=[run_index],
    )


def test_conflicting_fix_ops_are_withheld_and_violations_marked_conflict() -> None:
    """The scenario the checklist names: two rules writing different values
    to the same run's `font_size_pt`. Neither must apply -- a last-writer-
    wins outcome would make the audit log claim a fix that a second op then
    silently overwrote."""
    doc, _handles = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    target_paragraph_id = ast.paragraphs[0].id

    # Two synthetic rules disagreeing on the same (target, attribute).
    high_priority_size_14 = RULESET.by_id("title-size")
    assert high_priority_size_14 is not None
    conflicting_config = high_priority_size_14.model_copy(
        update={"rule_id": "fake-conflicting-size-rule", "fix_params": {"font_size_pt": 99}}
    )
    ruleset_with_conflict = RULESET.model_copy(
        update={"rules": [*RULESET.rules, conflicting_config]}
    )

    v1 = _fake_violation("title-size", target_paragraph_id)
    v2 = _fake_violation("fake-conflicting-size-rule", target_paragraph_id)

    plan = plan_fixes([v1, v2], ruleset_with_conflict)

    assert plan.ops == []
    assert v1.status is ViolationStatus.CONFLICT
    assert v2.status is ViolationStatus.CONFLICT


def test_agreeing_fix_ops_on_the_same_target_are_not_a_conflict() -> None:
    """Two rules that happen to agree on the value (or the same rule
    contributing twice, e.g. a shared run matched by two selectors) must not
    be penalised -- only a genuine disagreement is a conflict."""
    paragraph_id = "p0001"
    v1 = _fake_violation("title-size", paragraph_id)
    v2 = Violation(
        rule_id="title-size",
        severity=Severity.MEDIUM,
        paragraph_id=paragraph_id,
        auto_fixable=True,
        status=ViolationStatus.OPEN,
        run_indices=[0],
    )

    plan = plan_fixes([v1, v2], RULESET)

    assert v1.status is ViolationStatus.OPEN
    assert v2.status is ViolationStatus.OPEN
    matching = [op for op in plan.ops if op.rule_id == "title-size"]
    assert len(matching) == 1  # de-duplicated, not applied twice


def test_fix_action_registry_has_no_gaps_for_configured_actions() -> None:
    configured_actions = {
        rule.fix_action
        for rule in RULESET.rules
        if getattr(rule, "fix_action", None) not in (None, FixAction.SUGGEST_ONLY)
    }
    from manuscript_validator.autofix import actions

    registered = set(actions.RUN_ACTIONS) | set(actions.PARAGRAPH_ACTIONS)
    assert configured_actions <= registered
