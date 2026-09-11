"""Task 5 exit criteria: all 14 section labels assignable, tested against the
compliant fixture, a scrambled-order fixture, and heading-variant text; plus
`section_confidence`/`section_source` provenance and the missing-section and
front-matter-only degradation contracts that close FR-11's biggest gap --
a manuscript missing a whole section otherwise yields no violation at all.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from lxml import etree

from fixtures.factory import build_compliant, violating
from manuscript_validator.models import Section, SectionSource
from manuscript_validator.models.ast import Paragraph, Run
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.rules.schema import Ruleset
from manuscript_validator.segmenter import generate_section_present_rules, segment
from manuscript_validator.segmenter.headings import (
    HEADING_SCORE_THRESHOLD,
    evaluate_heading,
    normalize_heading_text,
)
from manuscript_validator.segmenter.heuristics import resolve_front_matter

RULESET = load_ruleset("journal_v1")

_ALL_SECTIONS_EXCEPT_UNKNOWN = frozenset(s for s in Section if s is not Section.UNKNOWN)


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_and_segment(doc: Document, ruleset: Ruleset = RULESET) -> Any:
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    result = segment(ast, ruleset)
    return ast, result


def _heading_paragraph(text: str, *, bold: bool = True) -> Paragraph:
    p = Paragraph(id="p1", text=text)
    p.runs = [Run(text=text, bold=bold)]
    return p


# --------------------------------------------------------------------------
# headings.evaluate_heading / normalize_heading_text
# --------------------------------------------------------------------------


def test_normalize_strips_enumeration_colon_and_case() -> None:
    assert normalize_heading_text("1. Discussion:") == "discussion"
    assert normalize_heading_text("  DISCUSSION  ") == "discussion"


def test_exact_vocabulary_match_scores_as_heading() -> None:
    synonyms = {Section.DISCUSSION: ["discussion"]}
    match = evaluate_heading(_heading_paragraph("DISCUSSION"), synonyms)
    assert match.is_heading
    assert match.section is Section.DISCUSSION
    assert match.match_kind == "exact"
    assert match.confidence == 1.0


def test_token_set_match_is_order_independent() -> None:
    synonyms = {Section.METHODS_AND_MATERIALS: ["methods and materials"]}
    match = evaluate_heading(_heading_paragraph("AND METHODS MATERIALS"), synonyms)
    assert match.is_heading
    assert match.section is Section.METHODS_AND_MATERIALS
    assert match.match_kind == "token_set"


def test_fuzzy_match_only_fires_above_structural_threshold() -> None:
    synonyms = {Section.DISCUSSION: ["discussion"]}
    # Bold + uppercase + short + no terminal period clears the structural
    # threshold on its own, so the misspelling is eligible for fuzzy matching.
    match = evaluate_heading(_heading_paragraph("DISCUSION"), synonyms)
    assert match.is_heading
    assert match.section is Section.DISCUSSION
    assert match.match_kind == "fuzzy"

    # The same misspelling with no structural signal at all (not bold, mixed
    # case, long) must not be dragged in by fuzzy matching alone.
    prose = _heading_paragraph(
        "A discusion of these results follows in the next several paragraphs.",
        bold=False,
    )
    match2 = evaluate_heading(prose, synonyms)
    assert not match2.is_heading


def test_caption_like_prose_is_not_a_heading() -> None:
    synonyms = {Section.RESULT: ["result"]}
    caption = _heading_paragraph("Table I. Summary of primary outcomes.", bold=False)
    match = evaluate_heading(caption, synonyms)
    assert not match.is_heading


def test_empty_paragraph_scores_no_heading() -> None:
    match = evaluate_heading(Paragraph(id="p1", text="   "), {})
    assert not match.is_heading


def test_heading_with_no_vocabulary_hit_can_still_score_as_a_heading() -> None:
    # All-bold + uppercase alone clears the threshold even with an empty
    # vocabulary -- it is "a heading of unknown section", not "not a heading".
    match = evaluate_heading(_heading_paragraph("STUDY LIMITATIONS"), {})
    assert match.score >= HEADING_SCORE_THRESHOLD
    assert match.is_heading
    assert match.section is None


# --------------------------------------------------------------------------
# heuristics.resolve_front_matter
# --------------------------------------------------------------------------


def _run(text: str, *, superscript: bool = False) -> Run:
    return Run(text=text, superscript=superscript)


def test_resolve_front_matter_assigns_title_author_abstract_affiliation() -> None:
    title = Paragraph(id="p1", text="A Study Of Something Important")
    author = Paragraph(id="p2", text="Jane Doe, John Roe")
    author.runs = [_run("Jane Doe"), _run("1", superscript=True), _run(", John Roe")]
    abstract_heading = Paragraph(id="p3", text="Abstract")
    abstract_body = Paragraph(id="p4", text="Background: this study evaluates something.")
    keywords = Paragraph(id="p5", text="Keywords: a, b, c")
    affiliation = Paragraph(id="p6", text="Dept of Medicine. Email: a@b.com. ORCID: 0000.")

    synonyms = {Section.ABSTRACT: ["abstract"]}
    paragraphs = [title, author, abstract_heading, abstract_body, keywords, affiliation]
    assignments = resolve_front_matter(paragraphs, synonyms)

    assert assignments["p1"].section is Section.TITLE
    assert assignments["p2"].section is Section.AUTHOR
    assert assignments["p3"].section is Section.ABSTRACT
    assert assignments["p3"].source is SectionSource.HEADING
    assert assignments["p4"].section is Section.ABSTRACT
    assert assignments["p5"].section is Section.ABSTRACT
    assert assignments["p6"].section is Section.AFFILIATION


def test_resolve_front_matter_falls_back_positionally_without_signals() -> None:
    # No "Abstract" heading, no ORCID/email in the last paragraph: the state
    # machine still produces a full, ordered assignment via fallbacks.
    title = Paragraph(id="p1", text="A Study Of Something Important")
    author = Paragraph(id="p2", text="Jane Doe")
    body = Paragraph(id="p3", text="This study evaluates something over several words.")
    last = Paragraph(id="p4", text="Unit of Medicine, Some City")

    assignments = resolve_front_matter([title, author, body, last], {})
    assert assignments["p1"].section is Section.TITLE
    assert assignments["p2"].section is Section.AUTHOR
    assert assignments["p3"].section is Section.ABSTRACT
    assert assignments["p4"].section is Section.AFFILIATION


def test_resolve_front_matter_on_empty_input() -> None:
    assert resolve_front_matter([], {}) == {}


# --------------------------------------------------------------------------
# section_index.segment: compliant, scrambled-order, and heading-variant
# --------------------------------------------------------------------------


def test_segment_assigns_all_14_labels_on_the_compliant_fixture() -> None:
    doc, _handles = build_compliant()
    ast, result = _build_and_segment(doc)

    assert result.detected_sections == _ALL_SECTIONS_EXCEPT_UNKNOWN
    assert result.missing_sections == frozenset()
    assert not result.front_matter_only

    title = ast.paragraphs[0]
    assert title.section is Section.TITLE
    assert title.section_source is SectionSource.POSITIONAL


def test_segment_assigns_section_confidence_and_source_provenance() -> None:
    doc, _handles = build_compliant()
    ast, _result = _build_and_segment(doc)

    intro_heading = next(p for p in ast.paragraphs if p.text.strip() == "INTRODUCTION")
    assert intro_heading.section is Section.INTRODUCTION
    assert intro_heading.section_source is SectionSource.HEADING
    assert intro_heading.section_confidence == 1.0

    intro_body = next(p for p in ast.paragraphs if "established evidence" in p.text)
    assert intro_body.section is Section.INTRODUCTION
    assert intro_body.section_source is SectionSource.INHERITED


def test_segment_table_cell_paragraphs_inherit_the_enclosing_section() -> None:
    """A bold, short cell like "Metric" scores as a heading by every
    structural signal `headings.py` uses; table cells must never run through
    heading detection at all, only inherit."""
    doc, _handles = build_compliant()
    ast, _result = _build_and_segment(doc)

    cell_paragraphs = [p for p in ast.paragraphs if p.in_table is not None]
    assert cell_paragraphs
    for cell in cell_paragraphs:
        assert cell.section is Section.RESULT
        assert cell.section_source is SectionSource.INHERITED


def test_segment_is_order_independent_for_scrambled_sections() -> None:
    """The `doc-order` fixture swaps Conflict of Interest and Funding --
    segmentation must still label each correctly regardless of position;
    sequence checking is `doc-order`'s own (semantic) job, not the
    segmenter's."""
    doc, _handles = violating("doc-order")
    _ast, result = _build_and_segment(doc)

    assert Section.CONFLICT_OF_INTEREST in result.detected_sections
    assert Section.FUNDING in result.detected_sections
    assert result.missing_sections == frozenset()


def test_segment_recognises_a_synonym_heading_variant() -> None:
    doc, handles = build_compliant()
    handles.introduction_heading_run.text = "BACKGROUND"
    ast, result = _build_and_segment(doc)

    heading = next(p for p in ast.paragraphs if p.text.strip() == "BACKGROUND")
    assert heading.section is Section.INTRODUCTION
    assert heading.section_source is SectionSource.HEADING
    assert Section.INTRODUCTION in result.detected_sections


def test_segment_reports_a_missing_required_section() -> None:
    doc, _handles = build_compliant()
    body = doc.element.body
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    to_remove: list[Any] = []
    capturing = False
    for p in list(body.iterchildren()):
        if etree.QName(p).localname != "p":
            continue
        text = "".join(t.text or "" for t in p.iter(f"{ns}t"))
        if text.strip() == "DISCUSSION":
            capturing = True
        elif capturing and text.strip() == "CONCLUSION":
            break
        if capturing:
            to_remove.append(p)
    for p in to_remove:
        p.getparent().remove(p)

    _ast, result = _build_and_segment(doc)
    assert Section.DISCUSSION in result.missing_sections
    assert Section.DISCUSSION not in result.detected_sections


def test_segment_flags_front_matter_only_when_no_body_headings_are_found() -> None:
    """A document with front matter but nothing that scores as a body
    heading must degrade honestly rather than guess section boundaries."""
    doc = Document()
    doc.add_paragraph("A Study Of Something Important")
    doc.add_paragraph("Jane Doe")
    p = doc.add_paragraph()
    p.add_run("Abstract").italic = True
    doc.add_paragraph("Background: this study evaluates something over many words.")
    doc.add_paragraph("Dept of Medicine. Email: a@b.com. ORCID: 0000-0001.")
    # Prose that mentions section names but never scores as a heading: not
    # bold, not uppercase, ordinary sentence-length paragraphs.
    doc.add_paragraph("This introduction describes the background of the study in prose.")
    doc.add_paragraph("The methods and materials used here are described next, in prose.")

    _ast, result = _build_and_segment(doc)
    assert result.front_matter_only
    assert result.degraded_reason == "section boundaries undetected"


# --------------------------------------------------------------------------
# generate_section_present_rules
# --------------------------------------------------------------------------


def test_generate_section_present_rules_covers_every_required_section() -> None:
    rules = generate_section_present_rules(RULESET)
    rule_ids = {rule.rule_id for rule in rules}
    assert rule_ids == {f"section-present-{s.value}" for s in RULESET.required_sections}
    assert "section-present-acknowledgement" not in rule_ids


def test_generate_section_present_rules_are_document_scoped() -> None:
    from manuscript_validator.models.enums import NodeType, Operator

    rules = generate_section_present_rules(RULESET)
    for rule in rules:
        assert rule.selector.node_type is NodeType.DOCUMENT
        assert rule.condition.operator is Operator.IS_NON_EMPTY
        section_name = rule.rule_id.removeprefix("section-present-")
        assert rule.condition.attribute == f"section:{section_name}"
        assert not rule.auto_fixable
