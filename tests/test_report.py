"""Task 11 exit criteria: the JSON report states what was and wasn't
checked, and every violation that has anywhere to point ends up as a native
Word comment on a clone of the *original* -- never `corrected.docx`, never
the redline.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import jsonschema
import pytest
from docx import Document
from lxml import etree

from fixtures.factory import build_compliant, violating
from manuscript_validator.models.enums import ViolationStatus
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.parser.binding import build_element_index
from manuscript_validator.report import UNCHECKED_SCOPES, annotate_document, build_report
from manuscript_validator.report.annotate_docx import anchor_runs
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.segmenter import segment

RULESET = load_ruleset("journal_v1")
_REPORT_SCHEMA = json.loads(
    (Path(__file__).parent / "schemas" / "report.schema.json").read_text("utf-8")
)


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _remove_section_block(doc: Document, heading_text: str, next_heading_text: str) -> None:
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = doc.element.body
    to_remove = []
    capturing = False
    for p in list(body.iterchildren()):
        if etree.QName(p).localname != "p":
            continue
        text = "".join(t.text or "" for t in p.iter(f"{ns}t"))
        if text.strip() == heading_text:
            capturing = True
        elif capturing and text.strip() == next_heading_text:
            break
        if capturing:
            to_remove.append(p)
    for p in to_remove:
        p.getparent().remove(p)


def _validate(source_bytes: bytes):
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    return ast, segmentation, rules_engine.validate(ast, RULESET, segmentation)


# --------------------------------------------------------------------------
# build_report
# --------------------------------------------------------------------------


def test_build_report_states_the_unchecked_scopes() -> None:
    doc, _handles = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    ast, segmentation, violations = _validate(source_bytes)

    report = build_report(ast.document_id, RULESET.ruleset_version, violations, segmentation)

    assert list(report.unchecked_scopes) == list(UNCHECKED_SCOPES)
    assert report.document_id == ast.document_id
    assert report.ruleset_version == RULESET.ruleset_version


def test_build_report_lists_missing_sections_sorted() -> None:
    doc, _handles = build_compliant()
    _remove_section_block(doc, "DISCUSSION", "CONCLUSION")
    source_bytes = _roundtrip_bytes(doc)
    ast, segmentation, violations = _validate(source_bytes)

    report = build_report(ast.document_id, RULESET.ruleset_version, violations, segmentation)

    assert report.missing_sections == ["discussion"]


def test_build_report_output_validates_against_the_schema() -> None:
    doc, _handles = violating("heading-uppercase-literal")
    source_bytes = _roundtrip_bytes(doc)
    ast, segmentation, violations = _validate(source_bytes)

    report = build_report(ast.document_id, RULESET.ruleset_version, violations, segmentation)

    jsonschema.validate(json.loads(json.dumps(report.to_dict())), _REPORT_SCHEMA)


# --------------------------------------------------------------------------
# anchor_runs
# --------------------------------------------------------------------------


def test_anchor_runs_returns_existing_runs() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("hello")
    paragraph.add_run("world")

    runs = anchor_runs(paragraph._p)

    assert [r.text for r in runs] == ["hello", "world"]


def test_anchor_runs_falls_back_to_a_zero_width_run_for_an_empty_paragraph() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()

    runs = anchor_runs(paragraph._p)

    assert len(runs) == 1
    assert runs[0].text == ""


# --------------------------------------------------------------------------
# annotate_document
# --------------------------------------------------------------------------


def test_annotate_document_adds_a_comment_readable_after_reopening() -> None:
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    _ast, _segmentation, violations = _validate(source_bytes)

    annotate_doc = Document(BytesIO(source_bytes))
    index = build_element_index(annotate_doc)
    count = annotate_document(annotate_doc, index, violations)
    assert count >= 1

    reopened = Document(BytesIO(_roundtrip_bytes(annotate_doc)))
    comments = list(reopened.comments)
    assert any("Title must be font size 14 pt." in c.text for c in comments)
    assert all(c.author == "Manuscript Validator" for c in comments)


def test_annotate_document_anchors_a_missing_section_violation_via_anchor_paragraph_id() -> None:
    doc, _handles = build_compliant()
    _remove_section_block(doc, "DISCUSSION", "CONCLUSION")
    source_bytes = _roundtrip_bytes(doc)
    _ast, _segmentation, violations = _validate(source_bytes)

    missing = next(v for v in violations if v.rule_id == "section-present-discussion")
    assert missing.paragraph_id is None
    assert missing.anchor_paragraph_id is not None

    annotate_doc = Document(BytesIO(source_bytes))
    index = build_element_index(annotate_doc)
    count_before = count_after = 0
    count_before = sum(1 for v in violations if v.paragraph_id or v.anchor_paragraph_id)
    count_after = annotate_document(annotate_doc, index, violations)

    assert count_after == count_before
    reopened = Document(BytesIO(_roundtrip_bytes(annotate_doc)))
    assert any("discussion" in c.text for c in reopened.comments)


def test_annotate_document_skips_violations_with_no_anchor_at_all() -> None:
    doc, _handles = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    _ast, _segmentation, violations = _validate(source_bytes)

    semantic_violations = [v for v in violations if v.status is ViolationStatus.CHECK_FAILED]
    assert semantic_violations  # the compliant fixture still has semantic stubs
    unanchored = [
        v for v in semantic_violations if v.paragraph_id is None and v.anchor_paragraph_id is None
    ]

    annotate_doc = Document(BytesIO(source_bytes))
    index = build_element_index(annotate_doc)
    count = annotate_document(annotate_doc, index, unanchored)

    assert count == 0  # no exception, nothing added


def test_annotate_document_never_touches_the_ast_it_was_built_from() -> None:
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    ast, _segmentation, violations = _validate(source_bytes)
    ast_snapshot_before = ast.to_dict()

    annotate_doc = Document(BytesIO(source_bytes))
    index = build_element_index(annotate_doc)
    annotate_document(annotate_doc, index, violations)

    assert ast.to_dict() == ast_snapshot_before


@pytest.mark.parametrize("rule_id", [r.rule_id for r in RULESET.rules])
def test_annotate_document_does_not_crash_for_any_rule(rule_id: str) -> None:
    doc, _handles = violating(rule_id)
    source_bytes = _roundtrip_bytes(doc)
    _ast, _segmentation, violations = _validate(source_bytes)

    annotate_doc = Document(BytesIO(source_bytes))
    index = build_element_index(annotate_doc)
    annotate_document(annotate_doc, index, violations)  # must not raise

    Document(BytesIO(_roundtrip_bytes(annotate_doc)))  # must reopen cleanly
