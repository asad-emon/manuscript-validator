"""Task 8 exit criteria (action half): each fix action tested in isolation
against a synthetic node -- a single paragraph or run built directly with
python-docx, no fixture factory needed, per spec section 13.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from manuscript_validator.autofix import actions
from manuscript_validator.errors import FixApplicationError
from manuscript_validator.parser.effective import has_trailing_line_break


def _run(text: str, **font_kwargs):
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    for name, value in font_kwargs.items():
        setattr(run.font, name, value)
    return run._r


def _paragraph_with_runs(*texts: str):
    doc = Document()
    paragraph = doc.add_paragraph()
    for text in texts:
        paragraph.add_run(text)
    return paragraph._p


def test_set_font_clears_theme_refs_before_writing_explicit_font() -> None:
    """Spec-alignment risk 2: leaving a theme ref in place means Word keeps
    rendering the theme font even after `ascii`/`hAnsi` say otherwise."""
    run_el = _run("hello")
    r_fonts = run_el.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    r_fonts.set(qn("w:hAnsiTheme"), "minorHAnsi")

    before, after = actions.set_font(run_el, {"font_name": "Times New Roman"})

    assert after == {"font_name": "Times New Roman"}
    assert before == {"font_name": None}
    assert r_fonts.get(qn("w:asciiTheme")) is None
    assert r_fonts.get(qn("w:hAnsiTheme")) is None
    assert r_fonts.get(qn("w:ascii")) == "Times New Roman"


def test_set_font_size() -> None:
    run_el = _run("hello")
    before, after = actions.set_font_size(run_el, {"font_size_pt": 14})
    assert before == {"font_size_pt": None}
    assert after == {"font_size_pt": 14}
    assert actions._wrap(run_el).font.size == Pt(14)


def test_set_bold() -> None:
    run_el = _run("hello", bold=False)
    before, after = actions.set_bold(run_el, {"bold": True})
    assert before == {"bold": False}
    assert after == {"bold": True}
    assert actions._wrap(run_el).font.bold is True


def test_set_italic() -> None:
    run_el = _run("hello")
    before, after = actions.set_italic(run_el, {"italic": True})
    assert before == {"italic": None}
    assert after == {"italic": True}
    assert actions._wrap(run_el).font.italic is True


def test_set_superscript() -> None:
    run_el = _run("1")
    before, after = actions.set_superscript(run_el, {"superscript": True})
    assert before == {"superscript": None}
    assert after == {"superscript": True}
    assert actions._wrap(run_el).font.superscript is True


def test_set_uppercase_literal_uppercases_text_and_clears_caps() -> None:
    run_el = _run("introduction")
    run_el.get_or_add_rPr().get_or_add_caps()

    before, after = actions.set_uppercase_literal(run_el, {})

    assert before == {"text": "introduction"}
    assert after == {"text": "INTRODUCTION"}
    assert actions._wrap(run_el).text == "INTRODUCTION"
    assert actions._wrap(run_el).font.all_caps is False


def test_strip_trailing_colon_edits_the_last_visible_run() -> None:
    paragraph_el = _paragraph_with_runs("INTRODUCTION:")
    before, after = actions.strip_trailing_colon(paragraph_el, {})
    assert before == {"text": "INTRODUCTION:"}
    assert after == {"text": "INTRODUCTION"}


def test_strip_trailing_colon_skips_a_trailing_break_run() -> None:
    """The journal template's own line-break convention
    (`paragraph.add_run().add_break()`) must not be mistaken for the
    paragraph's real content."""
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("INTRODUCTION:")
    paragraph.add_run().add_break()

    before, after = actions.strip_trailing_colon(paragraph._p, {})
    assert before == {"text": "INTRODUCTION:"}
    assert after == {"text": "INTRODUCTION"}


def test_strip_trailing_colon_raises_with_no_visible_run() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break()
    with pytest.raises(FixApplicationError):
        actions.strip_trailing_colon(paragraph._p, {})


def test_insert_line_break_appends_a_break_run() -> None:
    paragraph_el = _paragraph_with_runs("INTRODUCTION")
    assert not has_trailing_line_break(paragraph_el)

    before, after = actions.insert_line_break(paragraph_el, {})

    assert before == {"line_break_after": False}
    assert after == {"line_break_after": True}
    assert has_trailing_line_break(paragraph_el)


def test_set_citation_brackets_converts_parens_to_square_brackets() -> None:
    run_el = _run("(1)")
    before, after = actions.set_citation_brackets(run_el, {})
    assert before == {"text": "(1)"}
    assert after == {"text": "[1]"}


def test_set_citation_brackets_raises_on_unrecognised_text() -> None:
    run_el = _run("not a citation")
    with pytest.raises(FixApplicationError):
        actions.set_citation_brackets(run_el, {})


def test_set_numbering_style_arabic_to_roman_preserves_value() -> None:
    paragraph_el = _paragraph_with_runs("Table 1. Summary of primary outcomes.")
    before, after = actions.set_numbering_style(paragraph_el, {"numbering_style": "roman"})
    assert before == {"text": "Table 1. Summary of primary outcomes."}
    assert after == {"text": "Table I. Summary of primary outcomes."}


def test_set_numbering_style_roman_to_arabic_preserves_value() -> None:
    paragraph_el = _paragraph_with_runs("Figure IV. Trend over time.")
    before, after = actions.set_numbering_style(paragraph_el, {"numbering_style": "arabic"})
    assert before == {"text": "Figure IV. Trend over time."}
    assert after == {"text": "Figure 4. Trend over time."}


def test_set_numbering_style_raises_when_no_caption_token_found() -> None:
    paragraph_el = _paragraph_with_runs("This paragraph is not a caption.")
    with pytest.raises(FixApplicationError):
        actions.set_numbering_style(paragraph_el, {"numbering_style": "roman"})


def test_registries_cover_every_active_fix_action() -> None:
    """`set_caption_position`/`set_title_case` are deliberately absent
    (disabled in v1: both rules that would use them are `auto_fixable:
    false`), so no rule ever plans a `FixOp` naming either action."""
    from manuscript_validator.models.enums import FixAction

    registered = set(actions.RUN_ACTIONS) | set(actions.PARAGRAPH_ACTIONS)
    disabled = {FixAction.SET_CAPTION_POSITION, FixAction.SET_TITLE_CASE, FixAction.SUGGEST_ONLY}
    assert registered == set(FixAction) - disabled
