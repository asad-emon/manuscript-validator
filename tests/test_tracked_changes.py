"""Task 10 exit criteria: every audit entry appears as a revision, the
produced XML validates against the WordprocessingML XSD (docs/decisions.md
C3's verified constraints), and a real converter (LibreOffice, standing in
for Word) round-trips the result without rejecting it.
"""

from __future__ import annotations

import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from fixtures.factory import MUTATIONS, build_compliant, violating
from manuscript_validator.autofix import plan_fixes
from manuscript_validator.output.tracked_changes import (
    REVISION_AUTHOR,
    RevisionIdAllocator,
    apply_fix_plan_as_revisions,
)
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.parser.binding import build_element_index
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.segmenter import segment

RULESET = load_ruleset("journal_v1")
AUTO_FIXABLE_RULE_IDS = [rule.rule_id for rule in RULESET.rules if rule.auto_fixable]

_SCHEMA_PATH = Path(__file__).parent / "schemas" / "ISO-IEC29500-4_2016" / "wml.xsd"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WML_SCHEMA = etree.XMLSchema(etree.parse(str(_SCHEMA_PATH)))


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_and_plan(rule_id: str):
    doc, _handles = violating(rule_id)
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    violations = rules_engine.validate(ast, RULESET, segmentation)
    plan = plan_fixes(violations, RULESET)
    return source_bytes, plan


def _assert_document_part_is_schema_valid(document: Document) -> None:
    xml_bytes = etree.tostring(document.element)
    root = etree.fromstring(xml_bytes)
    root.attrib.pop(f"{{{_MC_NS}}}Ignorable", None)
    assert _WML_SCHEMA.validate(etree.ElementTree(root)), _WML_SCHEMA.error_log


@pytest.mark.parametrize("rule_id", AUTO_FIXABLE_RULE_IDS)
def test_tracked_document_validates_against_the_wml_schema(rule_id: str) -> None:
    source_bytes, plan = _build_and_plan(rule_id)
    assert any(op.rule_id == rule_id for op in plan.ops)

    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    audit_log = apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    assert any(entry.rule_id == rule_id for entry in audit_log)
    _assert_document_part_is_schema_valid(live_doc)


# --------------------------------------------------------------------------
# Per-mode structural assertions
# --------------------------------------------------------------------------


def test_format_only_fix_uses_rprchange_as_the_last_rpr_child() -> None:
    source_bytes, plan = _build_and_plan("title-size")
    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    title_paragraph = live_doc.paragraphs[0]._p
    run_el = title_paragraph.xpath(".//w:r[not(ancestor::w:rPr)]")[0]
    rpr = run_el.find(qn("w:rPr"))
    assert rpr is not None
    assert rpr[-1].tag == qn("w:rPrChange")

    rpr_change = rpr[-1]
    assert rpr_change.get(qn("w:author")) == REVISION_AUTHOR
    assert rpr_change.get(qn("w:date")) == "2026-09-12T10:00:00Z"
    original_rpr = rpr_change.find(qn("w:rPr"))
    assert original_rpr is not None
    original_size = original_rpr.find(qn("w:sz"))
    # 11pt, half-points -- the fixture's mutated (pre-fix) value.
    assert original_size.get(qn("w:val")) == "22"

    # The run itself was actually mutated to the new value, not just annotated.
    new_size = rpr.find(qn("w:sz"))
    assert new_size.get(qn("w:val")) == "28"  # 14pt


def test_two_format_fixes_on_the_same_run_produce_exactly_one_rprchange() -> None:
    """Regression: `w:rPrChange` is a `ZeroOrOne` child of `w:rPr` -- at most
    one is ever legal. Found only against real manuscripts, where a single
    author-line run commonly fails both `author-bold` and `author-size` at
    once; every python-docx-generated fixture violates at most one rule per
    run, so this never surfaced against a fixture until forced here.
    """
    doc, handles = build_compliant()
    MUTATIONS["author-bold"](handles)
    MUTATIONS["author-size"](handles)
    source_bytes = _roundtrip_bytes(doc)
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    violations = rules_engine.validate(ast, RULESET, segmentation)
    plan = plan_fixes(violations, RULESET)
    assert {"author-bold", "author-size"} <= {op.rule_id for op in plan.ops}

    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    target_id = next(op.target_id for op in plan.ops if op.rule_id == "author-bold")
    run_el = index.run(target_id)
    rpr = run_el.find(qn("w:rPr"))
    assert len(rpr.findall(qn("w:rPrChange"))) == 1

    # The single change still records the *original* pre-any-fix values for
    # both properties, not an intermediate state from whichever op ran first.
    original_rpr = rpr.find(qn("w:rPrChange")).find(qn("w:rPr"))
    original_bold = original_rpr.find(qn("w:b"))
    assert original_bold is not None and original_bold.get(qn("w:val")) == "0"
    original_size = original_rpr.find(qn("w:sz"))
    assert original_size.get(qn("w:val")) == "24"  # 12pt, the fixture's mutated value

    _assert_document_part_is_schema_valid(live_doc)


def test_text_edit_wraps_old_run_in_del_and_new_run_in_ins() -> None:
    source_bytes, plan = _build_and_plan("heading-uppercase-literal")
    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    op = next(o for o in plan.ops if o.rule_id == "heading-uppercase-literal")
    # Looked up by id, not by paragraph.text: python-docx's own `.text`
    # reads `./w:r` directly and would find nothing once that run is wrapped
    # in `w:ins` -- exactly the "never re-validate the redline" trap this
    # module's docstring names, so the test must not fall into it either.
    intro_paragraph = index.paragraph(op.paragraph_id)
    apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    del_el = intro_paragraph.find(qn("w:del"))
    ins_el = intro_paragraph.find(qn("w:ins"))
    assert del_el is not None
    assert ins_el is not None

    del_run = del_el.find(qn("w:r"))
    assert del_run.find(qn("w:t")) is None  # renamed, not merely duplicated
    del_text_el = del_run.find(qn("w:delText"))
    assert del_text_el is not None
    assert del_text_el.text == "Introduction"

    ins_run = ins_el.find(qn("w:r"))
    assert ins_run.find(qn("w:t")).text == "INTRODUCTION"

    # Revision ids are distinct, not accidentally reused across del/ins.
    assert del_el.get(qn("w:id")) != ins_el.get(qn("w:id"))


def test_pure_insertion_wraps_only_the_new_run() -> None:
    source_bytes, plan = _build_and_plan("heading-line-break-after")
    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    intro_paragraph = next(p._p for p in live_doc.paragraphs if p.text.strip() == "INTRODUCTION")
    assert intro_paragraph.find(qn("w:del")) is None
    ins_el = intro_paragraph.find(qn("w:ins"))
    assert ins_el is not None
    assert ins_el.find(qn("w:r")).find(qn("w:br")) is not None


# --------------------------------------------------------------------------
# Revision id allocation
# --------------------------------------------------------------------------


def test_revision_ids_start_at_one_with_no_existing_revisions() -> None:
    doc = Document()
    ids = RevisionIdAllocator(doc)
    assert ids.next_id() == 1
    assert ids.next_id() == 2


def test_revision_ids_start_above_a_pre_existing_co_author_revision() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    ins_el = paragraph._p.makeelement(qn("w:ins"), {qn("w:id"): "7"})
    paragraph._p.append(ins_el)

    ids = RevisionIdAllocator(doc)
    assert ids.next_id() == 8


# --------------------------------------------------------------------------
# Real-converter round trip
# --------------------------------------------------------------------------


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.slow
def test_soffice_round_trips_a_tracked_changes_document(tmp_path: Path) -> None:
    source_bytes, plan = _build_and_plan("heading-uppercase-literal")
    live_doc = Document(BytesIO(source_bytes))
    index = build_element_index(live_doc)
    apply_fix_plan_as_revisions(live_doc, index, plan, "2026-09-12T10:00:00Z")

    in_dir = tmp_path / "in"
    out_dir = tmp_path / "out"
    in_dir.mkdir()
    out_dir.mkdir()
    tracked_path = in_dir / "tracked.docx"
    live_doc.save(tracked_path)

    command = [
        "soffice", "--headless", "--convert-to", "docx",
        "--outdir", str(out_dir), str(tracked_path),
    ]
    result = subprocess.run(command, capture_output=True, timeout=60, check=False)
    assert result.returncode == 0, result.stderr

    converted = out_dir / "tracked.docx"
    reopened = Document(converted)
    assert len(reopened.paragraphs) > 0
