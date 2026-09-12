"""Task 12 exit criteria (rule-agnostic half): `semantic.build_evaluator()`
extracts the right text per rule, caches on that text, and turns a verdict
into a violation (or nothing) -- all against a `FakeSemanticClient`, with no
network dependency and no `google.genai` import anywhere in this file.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

import pytest
from docx import Document

from fixtures.factory import build_compliant, violating
from manuscript_validator.models.enums import ViolationStatus
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.cache import SemanticCache
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.rules.schema import SemanticVerdict
from manuscript_validator.rules.semantic import (
    SemanticCheckResult,
    build_evaluator,
    evaluate_api_key_missing,
    evaluate_stub,
    unsubstituted_placeholders,
)
from manuscript_validator.segmenter import segment

RULESET = load_ruleset("journal_v1")
SEMANTIC_RULES = {rule.rule_id: rule for rule in RULESET.rules if rule.rule_id in (
    "doc-order",
    "abstract-structure",
    "abstract-keywords-count",
    "affiliation-content",
    "result-text-before-figure",
    "reference-style-vancouver-format",
    "corresponding-author-content",
)}


@dataclass
class FakeSemanticClient:
    """Returns one queued result per call, in order; records every prompt it
    was asked to check, so a test can assert the cache prevented a second
    call."""

    results: list[SemanticCheckResult] = field(default_factory=list)
    prompts: list[str] = field(default_factory=list)

    def check(self, prompt: str) -> SemanticCheckResult:
        self.prompts.append(prompt)
        return self.results.pop(0)

    def test_connection(self) -> bool:
        return True


def _ok(violation: bool, **kwargs) -> SemanticCheckResult:
    return SemanticCheckResult(
        ok=True, verdict=SemanticVerdict(violation=violation, **kwargs)
    )


def _build_ast(doc: Document):
    buf = BytesIO()
    doc.save(buf)
    source_bytes = buf.getvalue()
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segment(ast, RULESET)
    return ast


def _cache() -> SemanticCache:
    return SemanticCache(ruleset_version=RULESET.ruleset_version, model_id="test-model")


# --------------------------------------------------------------------------
# build_evaluator: verdict -> violation mapping
# --------------------------------------------------------------------------


def test_evaluator_returns_none_when_verdict_says_no_violation() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False)])
    evaluate = build_evaluator(ast, client, _cache())

    result = evaluate(SEMANTIC_RULES["abstract-structure"])

    assert result is None


def test_evaluator_returns_needs_review_when_verdict_says_violation() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(
        results=[_ok(True, explanation="out of order", suggested_fix="reorder")]
    )
    evaluate = build_evaluator(ast, client, _cache())

    result = evaluate(SEMANTIC_RULES["abstract-structure"])

    assert result is not None
    assert result.status is ViolationStatus.NEEDS_REVIEW
    assert result.rule_id == "abstract-structure"
    assert result.explanation == "out of order"
    assert result.suggested_fix == "reorder"
    assert result.auto_fixable is False


def test_evaluator_returns_check_failed_when_client_fails() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(
        results=[SemanticCheckResult(ok=False, failure_reason="quota_exceeded")]
    )
    evaluate = build_evaluator(ast, client, _cache())

    result = evaluate(SEMANTIC_RULES["abstract-structure"])

    assert result is not None
    assert result.status is ViolationStatus.CHECK_FAILED
    assert result.failure_reason == "quota_exceeded"


def test_evaluator_check_failed_when_section_text_is_unavailable() -> None:
    """A section the segmenter never found (e.g. no corresponding-author
    block) has no text to send -- must degrade honestly, not call the client
    with an empty prompt."""
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    for paragraph in ast.paragraphs:
        section = paragraph.section
        if section is not None and section.value == "corresponding_author_address":
            paragraph.section = None
    client = FakeSemanticClient(results=[])
    evaluate = build_evaluator(ast, client, _cache())

    result = evaluate(SEMANTIC_RULES["corresponding-author-content"])

    assert result is not None
    assert result.status is ViolationStatus.CHECK_FAILED
    assert result.failure_reason == "section_text_unavailable"
    assert client.prompts == []  # never called


# --------------------------------------------------------------------------
# Caching
# --------------------------------------------------------------------------


def test_evaluator_hits_the_cache_on_a_second_identical_call() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False)])
    cache = _cache()
    evaluate = build_evaluator(ast, client, cache)

    first = evaluate(SEMANTIC_RULES["abstract-structure"])
    second = evaluate(SEMANTIC_RULES["abstract-structure"])

    assert first is None
    assert second is None
    assert len(client.prompts) == 1  # the second call was served from cache


def test_cache_key_changes_when_the_text_changes() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False), _ok(True, explanation="now different")])
    cache = _cache()
    evaluate = build_evaluator(ast, client, cache)
    evaluate(SEMANTIC_RULES["abstract-structure"])

    doc2, _handles2 = violating("abstract-structure")
    ast2 = _build_ast(doc2)
    evaluate2 = build_evaluator(ast2, client, cache)
    result2 = evaluate2(SEMANTIC_RULES["abstract-structure"])

    assert len(client.prompts) == 2  # different text, no cache hit
    assert result2 is not None


# --------------------------------------------------------------------------
# Text extraction, one assertion per semantic rule
# --------------------------------------------------------------------------


@pytest.mark.parametrize("rule_id", list(SEMANTIC_RULES))
def test_every_semantic_rule_extracts_non_empty_text_on_the_compliant_fixture(rule_id: str) -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False)])
    evaluate = build_evaluator(ast, client, _cache())

    evaluate(SEMANTIC_RULES[rule_id])

    assert len(client.prompts) == 1
    assert client.prompts[0].strip()
    assert not unsubstituted_placeholders(client.prompts[0])


def test_doc_order_extracts_the_detected_section_sequence() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False)])
    evaluate = build_evaluator(ast, client, _cache())

    evaluate(SEMANTIC_RULES["doc-order"])

    prompt = client.prompts[0]
    assert "title -> author -> abstract" in prompt
    assert "references" in prompt


def test_abstract_keywords_count_extracts_only_the_keywords_line() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    client = FakeSemanticClient(results=[_ok(False)])
    evaluate = build_evaluator(ast, client, _cache())

    evaluate(SEMANTIC_RULES["abstract-keywords-count"])

    assert "Keywords:" in client.prompts[0]
    assert "Background:" not in client.prompts[0]  # not the whole abstract


# --------------------------------------------------------------------------
# engine.validate() wiring
# --------------------------------------------------------------------------


def test_semantic_rules_run_in_parallel_not_serially() -> None:
    """Seven rules at ~3s serial is 20s of dead UI (the checklist's own
    framing) -- this proves `engine.validate()` actually parallelizes rather
    than merely accepting a `max_workers` constant that goes unused."""
    import time

    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    segmentation = segment(ast, RULESET)
    delay = 0.2

    def slow_evaluate(rule):
        time.sleep(delay)
        return None

    started = time.monotonic()
    rules_engine.validate(ast, RULESET, segmentation, evaluate_semantic=slow_evaluate)
    elapsed = time.monotonic() - started

    serial_time = delay * len(SEMANTIC_RULES)
    assert elapsed < serial_time * 0.75


def test_engine_uses_the_supplied_semantic_evaluator_instead_of_the_stub() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    segmentation = segment(ast, RULESET)
    client = FakeSemanticClient(results=[_ok(False)] * len(SEMANTIC_RULES))
    evaluate = build_evaluator(ast, client, _cache())

    violations = rules_engine.validate(ast, RULESET, segmentation, evaluate_semantic=evaluate)

    stub_reasons = {v.failure_reason for v in violations if v.rule_id in SEMANTIC_RULES}
    assert "semantic_check_not_implemented" not in stub_reasons


def test_engine_defaults_to_the_stub_when_no_evaluator_is_supplied() -> None:
    doc, _handles = build_compliant()
    ast = _build_ast(doc)
    segmentation = segment(ast, RULESET)

    violations = rules_engine.validate(ast, RULESET, segmentation)

    semantic_violations = [v for v in violations if v.rule_id in SEMANTIC_RULES]
    assert semantic_violations
    assert all(v.failure_reason == "semantic_check_not_implemented" for v in semantic_violations)


def test_evaluate_stub_matches_the_default_shape() -> None:
    stub_result = evaluate_stub(SEMANTIC_RULES["doc-order"])
    assert stub_result.status is ViolationStatus.CHECK_FAILED
    assert stub_result.auto_fixable is False


def test_evaluate_api_key_missing_is_distinct_from_the_not_implemented_stub() -> None:
    """Spec section 4.2: no configured key must surface as its own reason,
    not the generic `--no-semantic`/pre-Task-12 stub reason."""
    result = evaluate_api_key_missing(SEMANTIC_RULES["doc-order"])
    assert result.status is ViolationStatus.CHECK_FAILED
    assert result.failure_reason == "api_key_missing"
    assert result.failure_reason != evaluate_stub(SEMANTIC_RULES["doc-order"]).failure_reason
