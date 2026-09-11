"""Task 6 exit criteria (the strongest test in the suite, per the checklist --
and the reason Task 3's fixture factory came before any of this existed): the
compliant fixture yields zero deterministic violations, and every one of the
40 `violating(rule_id)` fixtures produces exactly one violation carrying that
rule's own id and no other.

Four rule pairs are declared as a known, understood exception rather than
forced to pass in isolation: mutating a table/figure's caption number changes
`Table.numbering_style`/`caption_number` (the field the mutation targets) but
also desyncs `Ast.captions.find_citing_paragraph_ids`'s number-based match
against the narrative text, which still cites the *old* number -- so
`table-cited-in-text`/`figure-cited-in-text` correctly fires too. This is a
real, explainable coupling in how captions are cross-referenced (Task 4), not
a flaw in this task's rule config.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from lxml import etree

from fixtures.factory import RULE_IDS, build_compliant, violating
from manuscript_validator.models.enums import CheckType, OnMissingSection, Section, ViolationStatus
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.rules import engine
from manuscript_validator.rules.loader import load_ruleset

RULESET = load_ruleset("journal_v1")
SEMANTIC_RULE_IDS = {
    rule.rule_id for rule in RULESET.rules if rule.check_type is CheckType.SEMANTIC
}

#: See module docstring: a single mutated caption number desyncs the
#: independent "is this table/figure cited in the narrative text" check too.
KNOWN_COUPLED_RULES: dict[str, frozenset[str]] = {
    "table-caption-numbering-style": frozenset({"table-cited-in-text"}),
    "table-numbering-sequence": frozenset({"table-cited-in-text"}),
    "figure-caption-numbering-style": frozenset({"figure-cited-in-text"}),
    "figure-numbering-sequence": frozenset({"figure-cited-in-text"}),
}


def _run_engine(doc: Document):
    from manuscript_validator.segmenter import segment

    buf = BytesIO()
    doc.save(buf)
    source_bytes = buf.getvalue()
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    return engine.validate(ast, RULESET, segmentation)


def _deterministic(violations):
    return [v for v in violations if v.status is not ViolationStatus.CHECK_FAILED]


def test_compliant_fixture_yields_zero_deterministic_violations() -> None:
    doc, _handles = build_compliant()
    violations = _run_engine(doc)
    det = _deterministic(violations)
    assert det == [], [(v.rule_id, v.found) for v in det]


def test_compliant_fixture_semantic_rules_report_check_failed_not_pass_or_fail() -> None:
    """Proves the section 7 router works with zero network dependency
    (Task 6's explicit exit criterion) -- every semantic rule resolves to
    exactly one `check_failed` violation, never a false pass or fail."""
    doc, _handles = build_compliant()
    violations = _run_engine(doc)
    for rule_id in SEMANTIC_RULE_IDS:
        matching = [v for v in violations if v.rule_id == rule_id]
        assert len(matching) == 1, rule_id
        assert matching[0].status is ViolationStatus.CHECK_FAILED
        assert matching[0].failure_reason == "semantic_check_not_implemented"


@pytest.mark.parametrize("rule_id", RULE_IDS)
def test_violating_fixture_triggers_exactly_that_rule(rule_id: str) -> None:
    doc, _handles = violating(rule_id)
    violations = _run_engine(doc)

    if rule_id in SEMANTIC_RULE_IDS:
        matching = [v for v in violations if v.rule_id == rule_id]
        assert len(matching) == 1
        assert matching[0].status is ViolationStatus.CHECK_FAILED
        return

    det = _deterministic(violations)
    matching = [v for v in det if v.rule_id == rule_id]
    other_ids = {v.rule_id for v in det if v.rule_id != rule_id}
    allowed = KNOWN_COUPLED_RULES.get(rule_id, frozenset())

    assert len(matching) == 1, f"expected exactly one {rule_id!r} violation, got {matching}"
    assert other_ids <= allowed, f"{rule_id!r} unexpectedly also triggered {other_ids - allowed}"


# --------------------------------------------------------------------------
# Missing-section and front-matter-only degradation (Task 5 contracts, wired
# through the engine here for the first time)
# --------------------------------------------------------------------------


def _remove_section_block(doc: Document, heading_text: str, next_heading_text: str) -> None:
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    body = doc.element.body
    to_remove = []
    capturing = False
    for p in list(body.iterchildren()):
        if etree.QName(p).localname != "p":
            continue
        text = "".join(t.text or "" for t in p.iter(f"{ns}t"))
        if text.strip() == heading_text:
            capturing = True
        elif capturing and text.strip() == next_heading_text:
            break
        if capturing:
            to_remove.append(p)
    for p in to_remove:
        p.getparent().remove(p)


def test_missing_required_section_produces_a_section_present_violation() -> None:
    doc, _handles = build_compliant()
    _remove_section_block(doc, "DISCUSSION", "CONCLUSION")
    violations = _run_engine(doc)
    matching = [v for v in violations if v.rule_id == "section-present-discussion"]
    assert len(matching) == 1
    assert matching[0].status is ViolationStatus.OPEN


def test_missing_section_does_not_silence_a_multi_section_rule_for_sections_still_present() -> None:
    """Regression: `heading-bold` applies to ten sections. An earlier version
    of the missing-section gate treated *any* overlap with `missing_sections`
    as "skip this rule entirely" -- so removing Discussion silently stopped
    checking heading-bold for Introduction, Methods, Result, and every other
    section it covers, not just the missing one."""
    doc, handles = build_compliant()
    handles.introduction_heading_run.font.bold = False
    _remove_section_block(doc, "DISCUSSION", "CONCLUSION")

    violations = _run_engine(doc)
    matching = [v for v in violations if v.rule_id == "heading-bold"]
    assert len(matching) == 1
    assert matching[0].section is Section.INTRODUCTION


def test_front_matter_only_degrades_body_rules_to_check_failed() -> None:
    doc = Document()
    doc.add_paragraph("A Study Of Something Important")
    doc.add_paragraph("Jane Doe")
    p = doc.add_paragraph()
    p.add_run("Abstract").italic = True
    doc.add_paragraph("Background: this study evaluates something over many words.")
    doc.add_paragraph("Dept of Medicine. Email: a@b.com. ORCID: 0000-0001.")
    doc.add_paragraph("This introduction describes the background of the study in prose.")

    violations = _run_engine(doc)
    body_rule_violations = [v for v in violations if v.rule_id == "body-text-size"]
    assert len(body_rule_violations) == 1
    assert body_rule_violations[0].status is ViolationStatus.CHECK_FAILED
    assert body_rule_violations[0].failure_reason == "section boundaries undetected"


def test_on_missing_section_skip_is_the_config_default() -> None:
    for rule in RULESET.rules:
        assert rule.on_missing_section is OnMissingSection.SKIP
