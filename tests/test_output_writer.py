"""Task 9 exit criteria (the strongest gate in the project, per the
checklist): the *written* file -- read back from an actual path on disk, not
the in-memory `Document` the writer just saved -- must re-parse, re-segment,
and re-validate to zero auto-fixable deterministic violations. In-memory
re-validation alone would miss serialization loss; only reopening the bytes
that actually hit disk proves FR-11 end to end.
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from fixtures.factory import build_compliant, violating
from manuscript_validator.autofix import plan_fixes
from manuscript_validator.models.enums import ViolationStatus
from manuscript_validator.models.report import utc_timestamp
from manuscript_validator.output import output_paths, write_bytes, write_corrected_document
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.segmenter import segment

RULESET = load_ruleset("journal_v1")
AUTO_FIXABLE_RULE_IDS = [rule.rule_id for rule in RULESET.rules if rule.auto_fixable]


def _roundtrip_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _validate(source_bytes: bytes):
    ast = build_ast(Document(BytesIO(source_bytes)), source_bytes)
    segmentation = segment(ast, RULESET)
    return ast, rules_engine.validate(ast, RULESET, segmentation)


@pytest.mark.parametrize("rule_id", AUTO_FIXABLE_RULE_IDS)
def test_corrected_file_on_disk_reparses_and_revalidates_clean(
    rule_id: str, tmp_path: Path
) -> None:
    doc, _handles = violating(rule_id)
    source_bytes = _roundtrip_bytes(doc)
    _ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)

    corrected_bytes, audit_log = write_corrected_document(
        source_bytes, plan, violations, utc_timestamp()
    )
    assert any(entry.rule_id == rule_id for entry in audit_log)

    source_path = tmp_path / "paper.docx"
    source_path.write_bytes(source_bytes)
    original_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()

    paths = output_paths(source_path, tmp_path)
    write_bytes(paths.corrected, corrected_bytes)

    # Reopen from the actual path, not the in-memory bytes just written --
    # this is what would catch a python-docx serialization bug that an
    # in-memory-only check could never see.
    reopened = Document(paths.corrected)
    written_bytes = paths.corrected.read_bytes()
    _ast2, violations2 = _validate(written_bytes)

    assert len(reopened.paragraphs) > 0  # reopens without corruption
    still_open = [
        v
        for v in violations2
        if v.rule_id == rule_id and v.status is not ViolationStatus.CHECK_FAILED
    ]
    assert still_open == []

    # FR-9: the original file on disk is untouched by any of the above.
    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == original_hash


def test_compliant_document_round_trips_with_an_empty_fix_plan(tmp_path: Path) -> None:
    doc, _handles = build_compliant()
    source_bytes = _roundtrip_bytes(doc)
    _ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)
    assert len(plan) == 0  # nothing to fix on an already-compliant document

    corrected_bytes, audit_log = write_corrected_document(
        source_bytes, plan, violations, utc_timestamp()
    )
    assert audit_log == []

    corrected_path = tmp_path / "paper.corrected.docx"
    write_bytes(corrected_path, corrected_bytes)
    reopened_ast, reopened_violations = _validate(corrected_path.read_bytes())

    det = [v for v in reopened_violations if v.status is not ViolationStatus.CHECK_FAILED]
    assert det == []
    assert len(reopened_ast.tables) == 2
    assert len(reopened_ast.figures) == 2


def test_corrected_document_preserves_structure_after_fixing_everything(tmp_path: Path) -> None:
    """Guards against regression toward C2's failure mode (docs/decisions.md):
    patching a clone must never lose images, tables, or paragraphs, the way
    regenerating a document from the AST projection would."""
    doc, _handles = build_compliant()
    # Break every auto-fixable rule at once by feeding the engine a document
    # that violates all of them, built by chaining every registered mutation.
    from fixtures.factory import MUTATIONS

    for rule_id in AUTO_FIXABLE_RULE_IDS:
        MUTATIONS[rule_id](_handles)
    source_bytes = _roundtrip_bytes(doc)
    original_ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)
    assert len(plan) > 0

    corrected_bytes, _audit_log = write_corrected_document(
        source_bytes, plan, violations, utc_timestamp()
    )
    corrected_path = tmp_path / "paper.corrected.docx"
    write_bytes(corrected_path, corrected_bytes)

    corrected_ast, corrected_violations = _validate(corrected_path.read_bytes())

    assert len(corrected_ast.tables) == len(original_ast.tables) == 2
    assert len(corrected_ast.figures) == len(original_ast.figures) == 2
    assert len(corrected_ast.paragraphs) == len(original_ast.paragraphs)

    still_open_ids = {
        v.rule_id for v in corrected_violations if v.status is not ViolationStatus.CHECK_FAILED
    }
    assert still_open_ids == set(), still_open_ids


def test_write_corrected_document_never_touches_the_callers_bytes(tmp_path: Path) -> None:
    doc, _handles = violating("title-size")
    source_bytes = _roundtrip_bytes(doc)
    source_path = tmp_path / "paper.docx"
    source_path.write_bytes(source_bytes)
    original_hash = hashlib.sha256(source_bytes).hexdigest()

    _ast, violations = _validate(source_bytes)
    plan = plan_fixes(violations, RULESET)
    write_corrected_document(source_bytes, plan, violations, utc_timestamp())

    assert hashlib.sha256(source_bytes).hexdigest() == original_hash
    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == original_hash
