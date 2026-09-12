"""Qt-free orchestration facade (Task 13).

Spec section 4 has the main window calling the pipeline stages directly, which
puts orchestration, error handling, and progress reporting inside a Qt class --
precisely the code that cannot be tested without a display and cannot be reused
by the CLI. This module is that code instead, and both shells call it.

Progress is a plain callable, not a Qt signal; `ui.workers` adapts it.
"""

from __future__ import annotations

import json
from collections.abc import Callable
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from docx import Document

from manuscript_validator.autofix import plan_fixes
from manuscript_validator.models.ast import Ast
from manuscript_validator.models.enums import ViolationStatus
from manuscript_validator.models.fix_plan import FixPlan
from manuscript_validator.models.report import ValidationReport, utc_timestamp
from manuscript_validator.models.violation import Violation
from manuscript_validator.output import (
    apply_fix_plan_as_revisions,
    output_paths,
    write_bytes,
    write_corrected_document,
)
from manuscript_validator.parser.ast_builder import build_ast
from manuscript_validator.parser.binding import build_element_index
from manuscript_validator.report import annotate_document, build_report
from manuscript_validator.rules import engine as rules_engine
from manuscript_validator.rules import semantic
from manuscript_validator.rules.cache import SemanticCache
from manuscript_validator.rules.llm_client import GeminiSemanticClient
from manuscript_validator.rules.loader import load_ruleset
from manuscript_validator.rules.schema import Ruleset, SemanticRule
from manuscript_validator.segmenter import segment

ProgressCallback = Callable[[str, float], None]


@dataclass(frozen=True)
class PipelineOptions:
    """Knobs the UI and CLI both expose.

    `api_key` is a plain value, not a lookup: *where* it comes from (an
    environment variable for the CLI, `config.settings_store` for the UI) is
    each shell's own business, not the pipeline's -- keeping this Qt-free
    facade from needing to know DPAPI exists.
    """

    output_dir: Path
    run_semantic: bool = True
    apply_fixes: bool = True
    ruleset_id: str = "journal_v1"
    api_key: str | None = None


@dataclass
class PipelineResult:
    """Everything a caller needs to render results and write outputs."""

    source: Path
    outputs: dict[str, Path] = field(default_factory=dict)
    violation_count: int = 0
    fixed_count: int = 0
    needs_review_count: int = 0
    check_failed_count: int = 0
    report: ValidationReport | None = None

    @property
    def max_open_severity(self) -> str | None:
        """The highest severity among violations still needing a human's
        attention -- `None` when there is nothing left to review."""
        if self.report is None:
            return None
        order = {"high": 2, "medium": 1, "low": 0}
        open_severities = [v.severity.value for v in self.report.violations if v.needs_review]
        if not open_severities:
            return None
        return max(open_severities, key=lambda s: order.get(s, 0))


def _select_semantic_evaluator(
    ast: Ast, ruleset: Ruleset, options: PipelineOptions
) -> Callable[[SemanticRule], Violation | None]:
    if not options.run_semantic:
        return semantic.evaluate_stub
    if not options.api_key:
        return semantic.evaluate_api_key_missing
    client = GeminiSemanticClient(api_key=options.api_key)
    cache = SemanticCache(ruleset_version=ruleset.ruleset_version, model_id=client.model_id)
    return semantic.build_evaluator(ast, client, cache)


def run(
    source: Path,
    options: PipelineOptions,
    progress: ProgressCallback | None = None,
) -> PipelineResult:
    """Run the full validate -> fix -> report pipeline.

    The source file is read to bytes once here and the path is never passed
    downstream (only `source` itself, for naming outputs by its stem), which
    makes FR-9 ("never mutates the original") a property of the design rather
    than a convention each stage has to honour.
    """

    def report_progress(stage: str, fraction: float) -> None:
        if progress is not None:
            progress(stage, fraction)

    report_progress("reading", 0.0)
    original_bytes = source.read_bytes()
    ruleset = load_ruleset(options.ruleset_id)

    report_progress("parsing", 0.1)
    ast = build_ast(Document(BytesIO(original_bytes)), original_bytes)

    report_progress("segmenting", 0.2)
    segmentation = segment(ast, ruleset)

    report_progress("validating", 0.3)
    evaluate_semantic = _select_semantic_evaluator(ast, ruleset, options)
    violations = rules_engine.validate(ast, ruleset, segmentation, evaluate_semantic)

    report_progress("planning fixes", 0.5)
    timestamp = utc_timestamp()
    fix_plan = plan_fixes(violations, ruleset) if options.apply_fixes else FixPlan()

    options.output_dir.mkdir(parents=True, exist_ok=True)
    paths = output_paths(source, options.output_dir)
    outputs: dict[str, Path] = {}

    report_progress("writing corrected document", 0.6)
    corrected_bytes, audit_log = write_corrected_document(
        original_bytes, fix_plan, violations, timestamp
    )
    write_bytes(paths.corrected, corrected_bytes)
    outputs["corrected"] = paths.corrected

    report_progress("writing tracked-changes redline", 0.7)
    tracked_document = Document(BytesIO(original_bytes))
    tracked_index = build_element_index(tracked_document)
    apply_fix_plan_as_revisions(tracked_document, tracked_index, fix_plan, timestamp)
    tracked_buffer = BytesIO()
    tracked_document.save(tracked_buffer)
    write_bytes(paths.tracked, tracked_buffer.getvalue())
    outputs["tracked"] = paths.tracked

    report_progress("annotating original", 0.8)
    annotated_document = Document(BytesIO(original_bytes))
    annotated_index = build_element_index(annotated_document)
    annotate_document(annotated_document, annotated_index, violations)
    annotated_buffer = BytesIO()
    annotated_document.save(annotated_buffer)
    write_bytes(paths.annotated, annotated_buffer.getvalue())
    outputs["annotated"] = paths.annotated

    report_progress("writing report and audit log", 0.9)
    validation_report = build_report(
        ast.document_id, ruleset.ruleset_version, violations, segmentation
    )
    paths.report.write_text(json.dumps(validation_report.to_dict(), indent=2), encoding="utf-8")
    outputs["report"] = paths.report

    audit_data = {"entries": [entry.to_dict() for entry in audit_log]}
    paths.audit.write_text(json.dumps(audit_data, indent=2), encoding="utf-8")
    outputs["audit"] = paths.audit

    report_progress("done", 1.0)

    return PipelineResult(
        source=source,
        outputs=outputs,
        violation_count=len(violations),
        fixed_count=sum(1 for v in violations if v.status is ViolationStatus.FIXED),
        needs_review_count=sum(1 for v in violations if v.needs_review),
        check_failed_count=sum(1 for v in violations if v.status is ViolationStatus.CHECK_FAILED),
        report=validation_report,
    )


__all__ = ["PipelineOptions", "PipelineResult", "ProgressCallback", "run"]
