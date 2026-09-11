"""`ElementIndex`: stable ids mapped to live XML elements.

Ids are positional at build time only; the index then holds element objects,
which survive reparenting (`tbl.addnext(cap_p)` preserves identity -- verified
in Task 3's structural fixture mutations). Because every clone is
byte-identical to the original at build time, ids resolve 1:1 across the
read, fixed, redline, and annotated clones -- `ids.walk_document` is the same
walk `ast_builder` uses, so nothing here can drift out of step with the AST.

Run splitting (Task 8) registers children as `p0042.r03/a`, `/b`, `/c`;
`ElementIndex.register_run` exists for that -- the initial build only ever
produces the plain `p0042.r03` form.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from docx.document import Document as DocumentObject

from manuscript_validator.errors import DocumentError
from manuscript_validator.parser.ids import ParagraphItem, iter_indexed_runs, walk_document


@dataclass
class ElementIndex:
    paragraphs: dict[str, Any] = field(default_factory=dict)
    runs: dict[str, Any] = field(default_factory=dict)
    tables: dict[str, Any] = field(default_factory=dict)

    def paragraph(self, paragraph_id: str) -> Any:
        try:
            return self.paragraphs[paragraph_id]
        except KeyError:
            raise DocumentError(f"no paragraph indexed for id {paragraph_id!r}") from None

    def run(self, run_id: str) -> Any:
        try:
            return self.runs[run_id]
        except KeyError:
            raise DocumentError(f"no run indexed for id {run_id!r}") from None

    def table(self, table_id: str) -> Any:
        try:
            return self.tables[table_id]
        except KeyError:
            raise DocumentError(f"no table indexed for id {table_id!r}") from None

    def register_run(self, run_id: str, run_element: Any) -> None:
        """Record a run produced by splitting an existing one (Task 8)."""
        self.runs[run_id] = run_element


def build_element_index(document: DocumentObject) -> ElementIndex:
    """Build a fresh index over `document`.

    Called once per clone (read, fixed, redline, annotated) -- never shared
    across clones, since each clone's elements are distinct Python/lxml
    objects even when byte-identical.
    """
    index = ElementIndex()
    for item in walk_document(document):
        if isinstance(item, ParagraphItem):
            index.paragraphs[item.id] = item.element
            for run_id, run_el in iter_indexed_runs(item.id, item.element):
                index.runs[run_id] = run_el
        else:
            index.tables[item.id] = item.element
    return index


__all__ = ["ElementIndex", "build_element_index"]
