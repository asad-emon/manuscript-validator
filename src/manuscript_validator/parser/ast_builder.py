"""Builds the AST by walking the document body in true order.

Delegates traversal and id assignment to `parser.ids.walk_document`, so
paragraph/run ids here are guaranteed to match `binding.build_element_index`
built against a byte-identical clone. Table-cell paragraphs land in the flat
`paragraphs` list tagged `in_table`/`cell`, as spec section 5.1 omits them
entirely and every run-level rule (`table-text-size`) needs to reach them.
"""

from __future__ import annotations

import hashlib
from typing import Any

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.text.run import Run as DocxRun

from manuscript_validator.models.ast import Ast, Figure, Paragraph, Run, Table
from manuscript_validator.parser import captions, shapes
from manuscript_validator.parser.effective import (
    EffectiveFormattingResolver,
    is_uppercase_literal,
    paragraph_text,
)
from manuscript_validator.parser.ids import (
    ParagraphItem,
    TableItem,
    iter_indexed_runs,
    walk_document,
)

_ALIGNMENT_NAMES = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _alignment_str(value: Any) -> str | None:
    if value is None:
        return None
    return _ALIGNMENT_NAMES.get(value, str(getattr(value, "name", value)).lower())


def _build_run(
    run_el: Any, docx_paragraph: DocxParagraph, resolver: EffectiveFormattingResolver
) -> Run:
    docx_run = DocxRun(run_el, docx_paragraph)
    text = docx_run.text
    name = resolver.resolve_font_name(docx_run, docx_paragraph)
    size = resolver.resolve_font_size(docx_run, docx_paragraph)
    bold = resolver.resolve_bold(docx_run, docx_paragraph)
    italic = resolver.resolve_italic(docx_run, docx_paragraph)
    superscript = resolver.resolve_superscript(docx_run, docx_paragraph)
    all_caps = resolver.resolve_all_caps(docx_run, docx_paragraph)
    return Run(
        text=text,
        font_name=name.value,
        font_name_source=name.source,
        font_size_pt=size.value,
        font_size_source=size.source,
        bold=bold.value,
        bold_source=bold.source,
        italic=italic.value,
        italic_source=italic.source,
        superscript=superscript.value,
        uppercase_literal=is_uppercase_literal(text),
        all_caps_property=bool(all_caps.value),
    )


def _build_paragraph(
    item: ParagraphItem, document: DocumentObject, resolver: EffectiveFormattingResolver
) -> Paragraph:
    docx_paragraph = DocxParagraph(item.element, document)
    style = docx_paragraph.style
    paragraph = Paragraph(
        id=item.id,
        text=paragraph_text(item.element),
        style_name=style.name or "" if style is not None else "",
        alignment=_alignment_str(docx_paragraph.alignment),
        in_table=item.in_table,
        cell=item.cell,
    )
    for _run_id, run_el in iter_indexed_runs(item.id, item.element):
        paragraph.runs.append(_build_run(run_el, docx_paragraph, resolver))
    return paragraph


def build_ast(document: DocumentObject, source_bytes: bytes) -> Ast:
    """Parse `document` into the read-only AST projection (docs/decisions.md C2).

    `source_bytes` is the exact bytes `document` was opened from -- used only
    to derive `document_id`; parsing never re-reads or mutates them (FR-9).
    """
    resolver = EffectiveFormattingResolver(document)
    ast = Ast(document_id=hashlib.sha256(source_bytes).hexdigest())

    # Keyed by the element itself, not `id(element)`: nothing else would keep
    # a bare id-of-object alive, and a garbage-collected lxml proxy's address
    # can be reused by an unrelated element, silently colliding two ids.
    element_to_paragraph_id: dict[Any, str] = {}
    table_items: list[TableItem] = []
    tables_by_id: dict[str, Table] = {}

    for item in walk_document(document):
        if isinstance(item, ParagraphItem):
            paragraph = _build_paragraph(item, document, resolver)
            ast.paragraphs.append(paragraph)
            element_to_paragraph_id[item.element] = item.id
            if item.in_table is not None:
                table = tables_by_id[item.in_table]
                table.cell_paragraph_ids.append(item.id)
                if table.cell_font_size_pt is None and paragraph.runs:
                    table.cell_font_size_pt = paragraph.runs[0].font_size_pt
        else:
            table = Table(id=item.id)
            ast.tables.append(table)
            tables_by_id[item.id] = table
            table_items.append(item)

    all_paragraph_texts = [(p.id, p.text) for p in ast.paragraphs]

    for item in table_items:
        table = tables_by_id[item.id]
        match = captions.find_table_caption(item.element)
        if match is None:
            continue
        table.caption_position = match.position
        table.numbering_style = match.numbering_style
        table.caption_number = match.number_token
        table.caption_paragraph_id = element_to_paragraph_id.get(match.element)
        others = [
            (pid, text) for pid, text in all_paragraph_texts if pid != table.caption_paragraph_id
        ]
        table.referenced_in_paragraph_ids = captions.find_citing_paragraph_ids(match, others)

    for i, image_p_el in enumerate(shapes.find_figure_paragraphs(document), start=1):
        figure_id = f"f{i:02d}"
        figure = Figure(
            id=figure_id, paragraph_id=element_to_paragraph_id.get(image_p_el)
        )
        match = captions.find_figure_caption(image_p_el)
        if match is not None:
            figure.caption_position = match.position
            figure.numbering_style = match.numbering_style
            figure.caption_number = match.number_token
            figure.caption_paragraph_id = element_to_paragraph_id.get(match.element)
            others = [
                (pid, text)
                for pid, text in all_paragraph_texts
                if pid != figure.caption_paragraph_id
            ]
            figure.referenced_in_paragraph_ids = captions.find_citing_paragraph_ids(match, others)
        ast.figures.append(figure)

    return ast


__all__ = ["build_ast"]
